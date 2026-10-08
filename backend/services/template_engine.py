import os
import re
from typing import Dict, Any, List, Set
from docx import Document

class TemplateEngine:
    def __init__(self):
        pass

    def extract_placeholders(self, docx_path: str) -> Set[str]:
        """
        Extracts all {{placeholder}} keys from a Word document.
        """
        if not os.path.exists(docx_path):
            return set()
        
        doc = Document(docx_path)
        placeholders = set()
        pattern = re.compile(r"\{\{([a-zA-Z0-9_\-\.]+)\}\}")

        for p in doc.paragraphs:
            for match in pattern.finditer(p.text):
                placeholders.add(match.group(1))

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for match in pattern.finditer(p.text):
                            placeholders.add(match.group(1))

        return placeholders

    def validate_template(self, docx_path: str, valid_keys: Set[str]) -> Dict[str, Any]:
        """
        Validates template placeholders against recognized canonical field keys.
        """
        found = self.extract_placeholders(docx_path)
        unknown = found - valid_keys - {"SECTION_CONTENT_BLOCK", "MASTER_NEWSLETTER_CONTENT", "issue_volume_issue", "issue_academic_year", "issue_semester", "issue_period", "issue_hod_message"}
        return {
            "valid": len(unknown) == 0,
            "found_placeholders": list(found),
            "unknown_placeholders": list(unknown)
        }
