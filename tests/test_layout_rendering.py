import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from docx import Document
from backend.services.docx_generator import DOCXGenerator, dedupe_entries, ensure_grid, compact
from backend.services.config_service import parse_form_script

ENTRY = {
    "entry_id": "E1", "category_id": "cat_x", "featured_flag": 0,
    "title": "Quantum Widgets", "author_or_person": "A. Author", "organization_or_venue": "Journal of Widgets",
    "event_date": "May 2026", "description": "Body text here.", "badge_tag": "Scopus Q1",
    "evidence_url": "https://example.org/paper", "photo_url": "", "photo_caption": "",
}


def block(type, tag, x, y, w, h, **extra):
    return {"id": f"{type}{x}{y}", "type": type, "tag": tag, "x": x, "y": y, "w": w, "h": h, "style": extra.pop("style", {}), **extra}


def cells(doc):
    """Every distinct cell of the first table, in reading order, as (row, col, text)."""
    table = doc.tables[0]
    seen, out = [], []
    for r, row in enumerate(table.rows):
        for c, cell in enumerate(row.cells):
            if any(cell._tc is t for t in seen):
                continue
            seen.append(cell._tc)
            if cell.text.strip():
                out.append((r + 1, c + 1, cell.text.strip()))
    return out


class LayoutRenderingTest(unittest.TestCase):
    def render(self, layout, entries=None):
        path = os.path.join(tempfile.mkdtemp(), "out.docx")
        DOCXGenerator().generate_section_docx({"section_id": "s", "name": "Research"}, [], entries or [ENTRY], path, layout=layout)
        return Document(path)

    def test_elements_land_in_the_cells_they_occupy_on_the_canvas(self):
        layout = [block("headline", "title", 1, 1, 6, 2), block("text", "description", 7, 1, 6, 4)]
        found = cells(self.render(layout))
        self.assertEqual(found, [(1, 1, "Quantum Widgets"), (1, 7, "Body text here.")])   # side by side

    def test_stacked_layout_puts_second_element_below(self):
        layout = [block("headline", "title", 1, 1, 12, 2), block("text", "description", 1, 3, 12, 3)]
        found = cells(self.render(layout))
        self.assertEqual([(r, c) for r, c, _ in found], [(1, 1), (3, 1)])

    def test_empty_elements_leave_no_gap(self):
        layout = [block("headline", "title", 1, 1, 12, 2), block("byline", "evidence_url_missing", 1, 3, 12, 1),
                  block("text", "description", 1, 4, 12, 3)]
        found = cells(self.render(layout))
        self.assertEqual([(r, c) for r, c, _ in found], [(1, 1), (3, 1)])   # description moved up under the headline

    def test_unselected_fields_are_not_printed(self):
        text = "\n".join(t for _, _, t in cells(self.render([block("headline", "title", 1, 1, 12, 2)])))
        self.assertNotIn("A. Author", text)
        self.assertNotIn("Body text here.", text)

    def test_style_from_layout_is_applied(self):
        layout = [block("headline", "title", 1, 1, 12, 2, style={"fontSize": "24px", "textColor": "#c65911", "textAlign": "center", "vAlign": "middle"})]
        p = self.render(layout).tables[0].cell(0, 0).paragraphs[0]
        self.assertEqual(p.runs[0].font.size.pt, 18)
        self.assertEqual(str(p.runs[0].font.color.rgb), "C65911")
        self.assertEqual(p.alignment, 1)

    def test_button_is_a_real_hyperlink_with_layout_text(self):
        layout = [block("link", "evidence_url", 1, 1, 6, 2, linkText="Read the paper", style={"buttonStyle": "accent"})]
        xml = self.render(layout).element.xml
        self.assertIn("Read the paper", xml)
        self.assertIn("<w:hyperlink", xml)

    def test_column_header_is_used_when_the_field_is_not_on_the_entry(self):
        layout = [block("text", "", 1, 1, 12, 2, column="Weird Question")]
        entry = dict(ENTRY)
        gen = DOCXGenerator()
        gen._raw_cache["R1"] = {"Weird Question": "raw answer"}
        entry["response_id"] = "R1"
        path = os.path.join(tempfile.mkdtemp(), "o.docx")
        gen.generate_section_docx({"section_id": "s", "name": "S"}, [], [entry], path, layout=layout)
        self.assertIn("raw answer", "\n".join(t for _, _, t in cells(Document(path))))

    def test_old_flat_layouts_are_stacked(self):
        grid = ensure_grid([{"type": "text", "tag": "title"}, {"type": "text", "tag": "description"}])
        self.assertEqual([(b["type"], b["x"], b["y"], b["w"]) for b in grid], [("headline", 1, 1, 12), ("text", 1, 4, 12)])

    def test_compact_floats_up(self):
        blocks = compact([block("text", "a", 1, 9, 12, 2), block("text", "b", 1, 1, 12, 2)])
        self.assertEqual(sorted(b["y"] for b in blocks), [1, 3])

    def test_duplicates_collapse(self):
        self.assertEqual(len(dedupe_entries([ENTRY, dict(ENTRY, entry_id="E2")])), 1)


class ScriptParsingTest(unittest.TestCase):
    SCRIPT = '''
    function a() {
      const form = FormApp.create("My Form");
      form.setTitle("ignored");
      form.addTextItem().setTitle("Full Name").setRequired(true);
      const cat = form.addMultipleChoiceItem();
      cat.setTitle("Topic").setChoiceValues(["One", "Two", 'Three']);
      form.addPageBreakItem().setTitle("Section 2");
      form.addParagraphTextItem().setTitle("Abstract");
    }
    function b() { const form = FormApp.create("Second"); form.addDateItem().setTitle("When"); }
    '''

    def test_questions_in_order_with_choices(self):
        forms = parse_form_script(self.SCRIPT)
        self.assertEqual([f["title"] for f in forms], ["My Form", "Second"])
        q = forms[0]["questions"]
        self.assertEqual([x["name"] for x in q], ["Full Name", "Topic", "Abstract"])
        self.assertTrue(q[0]["required"])
        self.assertEqual(q[1]["choices"], ["One", "Two", "Three"])
        self.assertEqual(forms[1]["questions"][0]["name"], "When")


if __name__ == "__main__":
    unittest.main()
