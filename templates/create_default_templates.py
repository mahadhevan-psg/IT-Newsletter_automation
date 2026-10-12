import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

TEMPLATES_DIR = os.path.dirname(__file__)

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top_dxa: int = 140, bottom_dxa: int = 140, left_dxa: int = 180, right_dxa: int = 180):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top_dxa}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom_dxa}" w:type="dxa"/>'
        f'<w:left w:w="{left_dxa}" w:type="dxa"/>'
        f'<w:right w:w="{right_dxa}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def create_section_template(filename: str, title: str, accent_hex: str = "1F4E78"):
    doc = Document()
    
    # 0.75-inch Page Margins
    for s in doc.sections:
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)

    # Section Top Masthead Banner
    header_table = doc.add_table(rows=1, cols=1)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = header_table.rows[0].cells[0]
    set_cell_background(cell, accent_hex)
    set_cell_margins(cell, top_dxa=160, bottom_dxa=160, left_dxa=200, right_dxa=200)
    cell.width = Inches(7.0)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_dept = p.add_run("IT CHRONICLES  |  DEPARTMENT OF INFORMATION TECHNOLOGY\n")
    run_dept.font.name = "Arial"
    run_dept.font.size = Pt(8.5)
    run_dept.font.bold = True
    run_dept.font.color.rgb = RGBColor(220, 235, 252)

    run = p.add_run(title.upper())
    run.font.name = "Arial"
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Placeholders Marker
    p_ph = doc.add_paragraph()
    r_ph = p_ph.add_run("{{SECTION_CONTENT_BLOCK}}")
    r_ph.font.name = "Arial"
    r_ph.font.size = Pt(10)
    r_ph.font.color.rgb = RGBColor(160, 160, 160)

    out_path = os.path.join(TEMPLATES_DIR, filename)
    doc.save(out_path)
    print(f"Created template: {out_path}")

def create_master_template():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)

    # ------------------ COVER PAGE ------------------
    # Top Dept Tag
    p_inst = doc.add_paragraph()
    p_inst.paragraph_format.space_before = Pt(36)
    p_inst.paragraph_format.space_after = Pt(4)
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run("DEPARTMENT OF INFORMATION TECHNOLOGY")
    r_inst.font.name = "Arial"
    r_inst.font.size = Pt(11)
    r_inst.font.bold = True
    r_inst.font.color.rgb = RGBColor(90, 107, 124)

    # Magazine Main Masthead
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(6)
    run_t = p_title.add_run("IT CHRONICLES")
    run_t.font.name = "Arial"
    run_t.font.size = Pt(38)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(31, 78, 120)

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(28)
    run_s = p_sub.add_run("The Biannual Journal of Innovation, Research & Student Excellence")
    run_s.font.name = "Arial"
    run_s.font.size = Pt(13)
    run_s.font.italic = True
    run_s.font.color.rgb = RGBColor(90, 107, 124)

    # Issue Info Ribbon Card
    info_table = doc.add_table(rows=1, cols=1)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_info = info_table.rows[0].cells[0]
    cell_info.width = Inches(7.0)
    set_cell_background(cell_info, "EBF3F9")
    set_cell_margins(cell_info, top_dxa=120, bottom_dxa=120, left_dxa=160, right_dxa=160)
    
    p_ribbon = cell_info.paragraphs[0]
    p_ribbon.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_rib = p_ribbon.add_run("{{issue_volume_issue}}   •   Academic Year: {{issue_academic_year}}   •   Semester: {{issue_semester}}   •   Period: {{issue_period}}")
    r_rib.font.name = "Arial"
    r_rib.font.size = Pt(9.5)
    r_rib.font.bold = True
    r_rib.font.color.rgb = RGBColor(31, 78, 120)

    doc.add_paragraph().paragraph_format.space_after = Pt(36)

    # Feature Preview Box
    prev_table = doc.add_table(rows=1, cols=1)
    prev_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_p = prev_table.rows[0].cells[0]
    cell_p.width = Inches(7.0)
    set_cell_background(cell_p, "F8FAFC")
    set_cell_margins(cell_p, top_dxa=160, bottom_dxa=160, left_dxa=180, right_dxa=180)

    p_in = cell_p.paragraphs[0]
    r_in = p_in.add_run("★  IN THIS ISSUE  ★\n")
    r_in.font.size = Pt(11)
    r_in.font.bold = True
    r_in.font.color.rgb = RGBColor(198, 89, 17)

    p_in_body = cell_p.add_paragraph()
    p_in_body.paragraph_format.line_spacing = 1.3
    bullets = [
        "◆ Research Milestones & Funded Grants: Break-through publications and national research grants.",
        "◆ Faculty Honors & Global Keynotes: National awards, expert symposiums, and certifications.",
        "◆ Student Hackathon Champions: National winners, placement stars, and creative tech innovations.",
        "◆ Industry & Outreach Collaborations: MoUs, industry visits, and alumni leadership spotlights."
    ]
    for b in bullets:
        r_b = p_in_body.add_run(b + "\n")
        r_b.font.size = Pt(9.5)
        r_b.font.color.rgb = RGBColor(50, 50, 50)

    doc.add_page_break()

    # ------------------ LEADERSHIP & EDITORIAL PAGE ------------------
    p_ed_title = doc.add_paragraph()
    p_ed_title.paragraph_format.space_before = Pt(16)
    p_ed_title.paragraph_format.space_after = Pt(10)
    r_ed = p_ed_title.add_run("MESSAGES & EDITORIAL LEADERSHIP")
    r_ed.font.name = "Arial"
    r_ed.font.size = Pt(16)
    r_ed.font.bold = True
    r_ed.font.color.rgb = RGBColor(31, 78, 120)

    # HOD Message Card
    hod_table = doc.add_table(rows=1, cols=1)
    hod_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_hod = hod_table.rows[0].cells[0]
    c_hod.width = Inches(7.0)
    set_cell_background(c_hod, "F8FAFC")
    set_cell_margins(c_hod, top_dxa=160, bottom_dxa=160, left_dxa=180, right_dxa=180)

    # Accent left border
    tcPr = c_hod._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="30" w:space="0" w:color="1F4E78"/>'
        f'<w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

    p_h_t = c_hod.paragraphs[0]
    r_h_t = p_h_t.add_run("Head of Department's Address\n")
    r_h_t.font.size = Pt(12)
    r_h_t.font.bold = True
    r_h_t.font.color.rgb = RGBColor(31, 78, 120)

    p_h_msg = c_hod.add_paragraph()
    p_h_msg.paragraph_format.line_spacing = 1.2
    r_h_msg = p_h_msg.add_run("“{{issue_hod_message}}”")
    r_h_msg.font.size = Pt(10)
    r_h_msg.font.italic = True
    r_h_msg.font.color.rgb = RGBColor(34, 34, 34)

    p_h_sig = c_hod.add_paragraph()
    p_h_sig.paragraph_format.space_before = Pt(8)
    r_h_sig = p_h_sig.add_run("— Dr. Department Head, Ph.D.\nProfessor & Head, Department of Information Technology")
    r_h_sig.font.size = Pt(9)
    r_h_sig.font.bold = True
    r_h_sig.font.color.rgb = RGBColor(90, 107, 124)

    doc.add_paragraph().paragraph_format.space_after = Pt(24)

    # Master Content Block
    doc.add_paragraph("{{MASTER_NEWSLETTER_CONTENT}}")

    out_path = os.path.join(TEMPLATES_DIR, "default_master_template.docx")
    doc.save(out_path)
    print(f"Created Master template: {out_path}")

if __name__ == "__main__":
    create_section_template("research_section_template.docx", "Research & Innovation", "1F4E78")
    create_section_template("faculty_section_template.docx", "Faculty Achievements", "2E75B6")
    create_section_template("student_section_template.docx", "Student Achievements", "3B7A57")
    create_section_template("industry_section_template.docx", "Industry & Alumni Outreach", "C65911")
    create_section_template("departmental_section_template.docx", "Departmental Highlights", "5B6770")
    create_master_template()
