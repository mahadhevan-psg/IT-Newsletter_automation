import os
import sys
import unittest
from docx import Document

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from database.db_manager import init_db, get_entries, update_entry_status
from backend.services.ingestion_service import IngestionService
from backend.services.docx_generator import DOCXGenerator
from backend.services.newsletter_builder import NewsletterBuilder

class TestMagazineGeneration(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        init_db()
        cls.ingestion = IngestionService()
        cls.docx_gen = DOCXGenerator()
        cls.builder = NewsletterBuilder()

        # Paths to demo images
        base_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
        cls.hackathon_img = os.path.join(base_dir, "output", "demo_images", "hackathon_win.jpg")
        cls.faculty_img = os.path.join(base_dir, "output", "demo_images", "prof_award.jpg")
        cls.mou_img = os.path.join(base_dir, "output", "demo_images", "mou_signing.jpg")

    def test_01_magazine_ingestion_and_field_mapping(self):
        # Ingest a landmark student hackathon achievement with photo & caption
        res = self.ingestion.process_single_submission(
            form_id="form_student",
            raw_data={
                "Timestamp": "2026-10-12 09:00:00",
                "Email Address": "rohan.sharma@ita.edu",
                "Student Achievement Category": "Hackathons & Coding Competitions",
                "Student Name(s)": "Rohan Sharma, Anita Roy, Kevin Patel",
                "Name of Competition / Award / Project / Event": "Global Smart Cities Hackathon Championship 2026",
                "Organizing Institute / Company / Body": "TechGiant Inc. & IEEE",
                "Position / Rank / Prize / CTC (e.g. 1st Place, Cash Prize of $1000)": "1st Prize ($5,000)",
                "Key Impact Badge / Metric": "1st Prize Champions",
                "Date of Event / Award": "October 2026",
                "Project Abstract / Achievement Details": "Built an autonomous drone dispatch platform for disaster relief operations.",
                "Photo Link (Google Drive / Cloud URL)": self.hackathon_img,
                "Photo Caption & Persons in Image": "Left to Right: Rohan Sharma, Anita Roy, Kevin Patel with 1st Prize trophy",
                "Certificate / Proof Link": "https://hackathon.org/proof",
                "Is this a Landmark / Milestone Achievement?": "Yes (Nominate for Spotlight Feature)"
            }
        )
        self.assertTrue(res["success"])
        entry_id = res["entry_id"]

        # Verify entry in DB
        entries = get_entries(issue_id="ISSUE_2026_V1I1", section_id="sec_student")
        saved = next((e for e in entries if e["entry_id"] == entry_id), None)
        self.assertIsNotNone(saved)
        self.assertEqual(saved["badge_tag"], "1st Prize Champions")
        self.assertEqual(saved["photo_caption"], "Left to Right: Rohan Sharma, Anita Roy, Kevin Patel with 1st Prize trophy")
        if not res.get("duplicate"):   # re-running against the same database finds the earlier entry
            self.assertTrue(saved["featured_flag"])

        # Approve for inclusion
        update_entry_status(entry_id, "Approved", featured_flag=True)

    def test_02_generate_magazine_section_docx(self):
        output_path = os.path.join(os.path.dirname(__file__), "..", "output", "Test_Magazine_Student_Section.docx")
        entries = get_entries(issue_id="ISSUE_2026_V1I1", section_id="sec_student", approval_status="Approved")
        subcats = [{"category_id": "cat_stu_hackathons", "name": "Hackathons & Coding Competitions"}]

        result_path = self.docx_gen.generate_section_docx(
            section={"section_id": "sec_student", "name": "Student Achievements"},
            subcategories=subcats,
            approved_entries=entries,
            output_path=output_path
        )
        self.assertTrue(os.path.exists(result_path))
        
        # Verify DOCX structure
        doc = Document(result_path)
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("STUDENT ACHIEVEMENTS", text)
        for e in entries:
            self.assertIn(e["title"], text)

    def test_03_generate_master_newsletter_docx(self):
        output_filename = "Test_Master_Magazine_Newsletter.docx"
        result = self.builder.assemble_newsletter(
            issue_id="ISSUE_2026_V1I1",
            output_filename=output_filename
        )
        self.assertTrue(result["success"])
        out_path = result["master_docx_path"]
        self.assertTrue(os.path.exists(out_path))

        doc = Document(out_path)
        self.assertGreater(len(doc.paragraphs), 10)
        print(f"\n[SUCCESS] Master Newsletter generated at: {out_path} with {len(doc.paragraphs)} paragraphs and {len(doc.tables)} tables.")

if __name__ == "__main__":
    unittest.main()
