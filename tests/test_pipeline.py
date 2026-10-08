import os
import sys
import unittest
from docx import Document

# Add parent directory to path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from database.db_manager import init_db, get_entries, update_entry_status, get_all_issues
from backend.services.ingestion_service import IngestionService
from backend.services.docx_generator import DOCXGenerator
from backend.services.newsletter_builder import NewsletterBuilder

class TestNewsletterPipeline(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        init_db()
        cls.ingestion = IngestionService()
        cls.docx_gen = DOCXGenerator()
        cls.builder = NewsletterBuilder()

    def test_01_ingestion_and_normalization(self):
        # 1. Ingest Research Entry
        res_research = self.ingestion.process_single_submission(
            form_id="form_research",
            raw_data={
                "Timestamp": "2026-10-08 10:00:00",
                "Email Address": "dr.smith@ita.edu",
                "Research Subcategory": "Journal Publications",
                "Title of Publication / Patent / Project / Innovation": "Advanced Quantum AI for Document Automation",
                "Authors / Investigators / Team Members": "Dr. John Smith, Dr. Alice Johnson",
                "Venue / Journal / Publisher / Funding Agency / Organization": "IEEE Transactions on Automation",
                "Date of Publication / Award / Grant": "October 2026",
                "Impact Factor / Indexing (e.g., Scopus, Web of Science, IEEE, Q1/Q2)": "Impact Factor 8.5 / Scopus Q1",
                "DOI / Patent Number / Grant Amount / URL Link": "https://doi.org/10.1109/TA.2026.001",
                "Abstract / Executive Summary / Highlights": "This paper presents a novel quantum architecture for normalizing dynamic newsletter components.",
                "Photo Link (Google Drive / Cloud URL)": "https://drive.google.com/photo1.jpg",
                "Evidence / Certificate Link (Google Drive / Document URL)": "https://drive.google.com/cert1.pdf"
            }
        )
        self.assertTrue(res_research["success"])
        self.assertIsNotNone(res_research["entry_id"])

        # 2. Ingest Faculty Entry
        res_faculty = self.ingestion.process_single_submission(
            form_id="form_faculty",
            raw_data={
                "Timestamp": "2026-10-08 10:15:00",
                "Email Address": "prof.varma@ita.edu",
                "Achievement Category": "Awards & Honors Received",
                "Faculty Full Name": "Prof. Rajesh Varma",
                "Title of Award / Certification / Keynote / Workshop": "Best Researcher Award 2026",
                "Organization / Host Institution / Awarding Body": "International Science Foundation",
                "Date of Event / Award / Completion": "September 2026",
                "Level (International / National / State / University)": "International",
                "Description & Key Highlights": "Awarded for exceptional contributions to distributed cloud computing architectures."
            }
        )
        self.assertTrue(res_faculty["success"])

        # 3. Ingest Student Entry
        res_student = self.ingestion.process_single_submission(
            form_id="form_student",
            raw_data={
                "Timestamp": "2026-10-08 10:30:00",
                "Email Address": "student1@ita.edu",
                "Student Achievement Category": "Hackathons & Coding Competitions",
                "Student Name(s)": "Rohan Sharma, Anita Roy",
                "Name of Competition / Award / Project / Event": "National AI Hackathon 2026",
                "Organizing Institute / Company / Body": "TechCorp India",
                "Position / Rank / Prize / CTC (e.g. 1st Place, Cash Prize of $1000)": "1st Place (Cash Prize $2000)",
                "Date of Event / Award": "October 2026",
                "Project Abstract / Achievement Details": "Built an automated newsletter builder platform using document transformation engines."
            }
        )
        self.assertTrue(res_student["success"])

    def test_02_editorial_workflow(self):
        entries = get_entries(issue_id="ISSUE_2026_V1I1")
        self.assertGreater(len(entries), 0)

        # Approve all entries and set first entry as featured spotlight
        for idx, entry in enumerate(entries):
            is_feat = (idx == 0)
            success = update_entry_status(
                entry_id=entry["entry_id"],
                status="Approved",
                featured_flag=is_feat,
                performer="TestEditor"
            )
            self.assertTrue(success)

        approved = get_entries(issue_id="ISSUE_2026_V1I1", approval_status="Approved")
        self.assertEqual(len(approved), len(entries))

    def test_03_section_docx_generation(self):
        approved_research = get_entries(issue_id="ISSUE_2026_V1I1", section_id="sec_research", approval_status="Approved")
        sec_info = {
            "section_id": "sec_research",
            "name": "Research & Innovation",
            "template_file": os.path.join(os.path.dirname(__file__), "..", "templates", "research_section_template.docx")
        }
        subcategories = [
            {"category_id": "cat_journal_pub", "name": "Journal Publications"}
        ]

        out_path = os.path.join(os.path.dirname(__file__), "..", "output", "Test_Section_Research.docx")
        generated_file = self.docx_gen.generate_section_docx(
            section=sec_info,
            subcategories=subcategories,
            approved_entries=approved_research,
            output_path=out_path
        )

        self.assertTrue(os.path.exists(generated_file))
        doc = Document(generated_file)
        self.assertGreater(len(doc.paragraphs) + len(doc.tables), 0)

    def test_04_master_newsletter_assembly(self):
        res = self.builder.assemble_newsletter(
            issue_id="ISSUE_2026_V1I1",
            output_filename="Test_Master_Newsletter.docx"
        )
        self.assertTrue(res["success"])
        self.assertTrue(os.path.exists(res["master_docx_path"]))
        
        doc = Document(res["master_docx_path"])
        text = "\n".join([p.text for p in doc.paragraphs])
        self.assertIn("IT CHRONICLES", text)

if __name__ == "__main__":
    unittest.main()
